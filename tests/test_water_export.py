"""Unit test — export rekap air minum (v2.37.5).

Laporan resmi PDF & Excel landscape, TTD Finance (Dibuat oleh) &
Kepala Cabang (Mengetahui), filter sama dengan daftar.

Diuji:
- generator Excel  : valid .xlsx, judul, header kolom, ringkasan, TTD,
                     print setup landscape;
- generator PDF    : valid %PDF, landscape (w > h), berisi judul & TTD;
- endpoint         : xlsx & pdf content-type/filename benar, filter
                     diteruskan, OB ter-scope.

Jalankan: python3 -m pytest tests/test_water_export.py -v
"""
import sys
import os
from io import BytesIO
from datetime import date

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))

import pytest
from flask import Flask

import modules.routes_water as rw
from modules.routes_water import register_water_routes
from modules.water_report import generate_water_report_excel, WaterReportPDF

META = {
    'from': date(2026, 9, 1), 'to': date(2026, 10, 1),
    'filters_text': 'Status Terverifikasi',
    'company': {'company_name': 'PT BESTPROFIT FUTURES', 'company_subtitle': 'Cabang Surabaya'},
    'head_name': 'Budi Santoso',
    'finance_name': 'Rina Wijaya',
}

ROWS = [
    {'id': 1, 'display_id': 'WTR-SBY-20260907-0001', 'ob_name': 'Febri',
     'purchase_date': date(2026, 9, 3), 'status': 'verified', 'remark': 'OK', 'note': '',
     'items': [{'drink_type': 'Galon', 'brand': 'AQUA', 'satuan': 'galon', 'quantity': 3}]},
    {'id': 2, 'display_id': 'WTR-SBY-20260907-0002', 'ob_name': 'Edwin',
     'purchase_date': date(2026, 9, 5), 'status': 'pending', 'remark': '', 'note': 'catatan uji',
     'items': [{'drink_type': 'Botol', 'brand': 'Le Minerale', 'satuan': 'dus', 'quantity': 2}]},
]


# ================================================================
# Generator Excel
# ================================================================
class TestExcel:
    def test_valid_xlsx(self):
        from openpyxl import load_workbook
        data = generate_water_report_excel(ROWS, META)
        wb = load_workbook(BytesIO(data))
        ws = wb.active
        assert ws.title == 'Rekap Air Minum'

    def test_judul_dan_periode(self):
        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(generate_water_report_excel(ROWS, META)))
        ws = wb.active
        assert 'AIR MINUM' in str(ws['A1'].value)
        assert '01/09/2026' in str(ws['A2'].value)

    def test_header_kolom_lengkap(self):
        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(generate_water_report_excel(ROWS, META)))
        ws = wb.active
        headers = [ws.cell(row=5, column=i).value for i in range(1, 10)]
        assert headers[1] == 'No. Dokumen' and 'Status' in headers and 'Remark' in str(headers)

    def test_baris_data_dan_ringkasan(self):
        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(generate_water_report_excel(ROWS, META)))
        ws = wb.active
        assert ws.cell(row=6, column=2).value == 'WTR-SBY-20260907-0001'
        ringkasan = [c.value for row in ws.iter_rows() for c in row
                     if isinstance(c.value, str) and c.value.startswith('RINGKASAN')]
        assert ringkasan and 'Total 2' in ringkasan[0] and 'Terverifikasi 1' in ringkasan[0]

    def test_ttd_kepala_dan_finance(self):
        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(generate_water_report_excel(ROWS, META)))
        ws = wb.active
        semua = [str(c.value) for row in ws.iter_rows() for c in row if c.value]
        assert 'BUDI SANTOSO' in semua and 'RINA WIJAYA' in semua
        assert 'Kepala Cabang' in semua and 'Finance' in semua

    def test_print_setup_landscape(self):
        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(generate_water_report_excel(ROWS, META)))
        ws = wb.active
        assert ws.page_setup.orientation == 'landscape'

    def test_ttd_fallback_label(self):
        from openpyxl import load_workbook
        meta = dict(META, head_name='', finance_name='')
        wb = load_workbook(BytesIO(generate_water_report_excel(ROWS, meta)))
        ws = wb.active
        semua = [str(c.value) for row in ws.iter_rows() for c in row if c.value]
        assert 'KEPALA CABANG' in semua and 'FINANCE' in semua


# ================================================================
# Generator PDF
# ================================================================
class TestPDF:
    def test_valid_pdf_magic(self):
        data = WaterReportPDF().generate(ROWS, META)
        assert data[:5] == b'%PDF-'

    def test_landscape(self):
        from fpdf import FPDF
        pdf = WaterReportPDF()
        pdf.generate(ROWS, META)
        assert pdf.pdf.w > pdf.pdf.h  # landscape: lebar > tinggi

    def test_konten_judul_ttd(self):
        from tests.pdf_text import pdf_text as _pdf_text
        data = WaterReportPDF().generate(ROWS, META)
        text = _pdf_text(data)
        assert 'AIR MINUM' in text
        assert 'BUDI SANTOSO' in text and 'RINA WIJAYA' in text
        assert 'Kepala Cabang' in text and 'Finance' in text
        assert 'WTR-SBY-20260907-0001' in text

    def test_multi_page_tetap_valid(self):
        many = [dict(ROWS[0], id=i, display_id=f'WTR-SBY-20260907-{i:04d}') for i in range(3, 80)]
        data = WaterReportPDF().generate(many, META)
        assert data[:5] == b'%PDF-'
        assert data.count(b'/Type /Page') >= 2 or b'/Pages' in data


# ================================================================
# Endpoint
# ================================================================
class FakeCursor:
    def __init__(self, store):
        self.store = store
        self.executed = []
        self._rows = []

    def execute(self, sql, params=None):
        self.executed.append((sql, params))
        if 'SELECT wp.*' in sql:
            self._rows = []
        else:
            self._rows = []

    def fetchall(self):
        return list(self._rows)

    def fetchone(self):
        return None

    def close(self):
        pass


class FakeConn:
    def __init__(self, store):
        self.cursors = [FakeCursor(store)]

    def cursor(self, dictionary=False):
        return self.cursors[-1]

    def close(self):
        pass


@pytest.fixture()
def client():
    app = Flask(__name__)
    app.secret_key = 'test-secret'
    app.config['TESTING'] = True
    register_water_routes(app)
    with app.test_client() as c:
        yield c


def _login(client, role='finance', name='Cynthia Permatasaya', branch='SBY'):
    with client.session_transaction() as s:
        s['user_name'] = name
        s['full_name'] = name
        s['user_role'] = role
        s['branch_code'] = branch


def test_export_xlsx_content_type(client, monkeypatch):
    monkeypatch.setattr(rw, 'get_db_connection', lambda: FakeConn({}))
    _login(client)
    r = client.get('/api/water/purchases/export?format=xlsx')
    assert r.status_code == 200
    assert 'spreadsheetml' in r.headers['Content-Type']
    assert 'attachment' in r.headers['Content-Disposition']
    assert '.xlsx' in r.headers['Content-Disposition']


def test_export_pdf_content_type(client, monkeypatch):
    monkeypatch.setattr(rw, 'get_db_connection', lambda: FakeConn({}))
    _login(client)
    r = client.get('/api/water/purchases/export?format=pdf')
    assert r.status_code == 200
    assert r.headers['Content-Type'] == 'application/pdf'
    assert '.pdf' in r.headers['Content-Disposition']
    assert r.data[:5] == b'%PDF-'


def test_export_filter_diteruskan(client, monkeypatch):
    conns = []
    monkeypatch.setattr(rw, 'get_db_connection', lambda: (conns.append(FakeConn({})) or conns[-1]))
    _login(client)
    r = client.get('/api/water/purchases/export?format=xlsx&from=2026-08-01&to=2026-08-31&status=pending&q=AQUA')
    assert r.status_code == 200
    sql = conns[0].cursors[0].executed[0][0]
    params = conns[0].cursors[0].executed[0][1]
    assert 'wp.status = %s' in sql and 'LIKE %s' in sql
    assert '2026-08-01' in [str(p) for p in params]
    assert '%AQUA%' in [str(p) for p in params]


def test_export_ob_ter_scope(client, monkeypatch):
    conns = []
    monkeypatch.setattr(rw, 'get_db_connection', lambda: (conns.append(FakeConn({})) or conns[-1]))
    _login(client, role='ob', name='Edwin')
    r = client.get('/api/water/purchases/export?format=xlsx')
    assert r.status_code == 200
    sql = conns[0].cursors[0].executed[0][0]
    assert 'wp.ob_name = %s' in sql
