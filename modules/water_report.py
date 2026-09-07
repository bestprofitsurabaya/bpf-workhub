"""Generator laporan resmi rekap pengajuan air minum (v2.37.5).

Dua format, dipanggil dari GET /api/water/purchases/export:
- generate_water_report_excel(rows, meta) → bytes .xlsx (openpyxl, landscape)
- generate_water_report_pdf(rows, meta)   → bytes .pdf  (fpdf2, landscape)

Desain mengikuti pola laporan resmi yang sudah ada:
- Excel: generate_appointment_report (header biru, border tipis, ringkasan)
- PDF  : BPFBasePDF (kop identitas perusahaan, DejaVuSans, footer halaman)

TTD default dua pihak: Kepala Cabang (Mengetahui) & Finance (Dibuat oleh).
Nama dari system_config: water_head_name & water_finance_name.
"""
from io import BytesIO

STATUS_LABEL = {'pending': 'Menunggu Verifikasi',
                'verified': 'Terverifikasi',
                'rejected': 'Ditolak'}


def _fmt_date(v):
    from datetime import datetime, date
    if isinstance(v, (datetime, date)):
        return v.strftime('%d/%m/%Y')
    s = str(v or '')
    return s[:10] if s else '-'


def _items_text(p):
    items = p.get('items') or []
    if not items:
        return '-'
    return '; '.join(f"{i.get('brand', '-')} {i.get('drink_type', '')} "
                     f"{i.get('quantity', '')} {i.get('satuan', '')}".strip()
                     for i in items)


def _totals(rows):
    t = {'pending': 0, 'verified': 0, 'rejected': 0, 'qty': 0}
    for r in rows:
        t[r.get('status', 'pending')] = t.get(r.get('status'), 0) + 1
        for it in (r.get('items') or []):
            t['qty'] += int(it.get('quantity') or 0)
    return t


def _signatures(meta):
    """(nama_kepala, nama_finance) — fallback label bila belum diset."""
    head = (meta.get('head_name') or '').strip() or 'Kepala Cabang'
    fin = (meta.get('finance_name') or '').strip() or 'Finance'
    return head, fin


# ============================================================
# EXCEL (openpyxl) — landscape A4
# ============================================================
def generate_water_report_excel(rows, meta):
    """rows = list dict pengajuan (dengan items); meta = dict filter & TTD."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

    head_name, fin_name = _signatures(meta)
    t = _totals(rows)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Rekap Air Minum'

    thin = Border(left=Side(style='thin'), right=Side(style='thin'),
                  top=Side(style='thin'), bottom=Side(style='thin'))
    header_fill = PatternFill(start_color='1D4ED8', end_color='1D4ED8', fill_type='solid')
    title_font = Font(name='Arial', bold=True, size=14, color='1E293B')
    subtitle_font = Font(name='Arial', size=10, color='475569')
    header_font = Font(name='Arial', bold=True, size=9, color='FFFFFF')
    normal_font = Font(name='Arial', size=9)
    bold_font = Font(name='Arial', bold=True, size=9)

    # A No | B No Dokumen | C Tanggal | D OB | E Rincian | F Total Qty | G Status | H Remark | I Catatan
    widths = {'A': 5, 'B': 22, 'C': 12, 'D': 18, 'E': 46, 'F': 10, 'G': 18, 'H': 34, 'I': 34}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    last_col = 'I'

    ws.merge_cells(f'A1:{last_col}1')
    ws['A1'] = 'LAPORAN REKAP PENGAJUAN AIR MINUM'
    ws['A1'].font = title_font
    ws['A1'].alignment = Alignment(horizontal='center')

    ws.merge_cells(f'A2:{last_col}2')
    ident = meta.get('company') or {}
    ws['A2'] = (f"{ident.get('company_name', 'PT BESTPROFIT FUTURES')} — "
                f"{ident.get('company_subtitle', 'Cabang')}"
                f"  •  Periode {_fmt_date(meta.get('from'))} s/d {_fmt_date(meta.get('to'))}")
    ws['A2'].font = subtitle_font
    ws['A2'].alignment = Alignment(horizontal='center')

    ws.merge_cells(f'A3:{last_col}3')
    ftxt = meta.get('filters_text') or ''
    ws['A3'] = f'Filter: {ftxt}' if ftxt else 'Filter: Semua status'
    ws['A3'].font = Font(name='Arial', italic=True, size=9, color='64748B')
    ws['A3'].alignment = Alignment(horizontal='center')

    headers = ['No', 'No. Dokumen', 'Tanggal', 'OB', 'Rincian Item', 'Total Qty',
               'Status', 'Remark Verifikasi', 'Catatan']
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=5, column=ci, value=h)
        c.font = header_font
        c.fill = header_fill
        c.border = thin
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    r0 = 6
    for i, p in enumerate(rows, 1):
        vals = [i, p.get('display_id', '-'), _fmt_date(p.get('purchase_date')),
                p.get('ob_name', '-'), _items_text(p),
                sum(int(it.get('quantity') or 0) for it in (p.get('items') or [])),
                STATUS_LABEL.get(p.get('status'), p.get('status', '-')),
                p.get('remark') or '-', p.get('note') or '-']
        for ci, v in enumerate(vals, 1):
            c = ws.cell(row=r0 + i - 1, column=ci, value=v)
            c.font = normal_font
            c.border = thin
            if ci in (1, 3, 6):
                c.alignment = Alignment(horizontal='center')
            if ci == 7:
                c.alignment = Alignment(horizontal='center')
            if ci in (5, 8, 9):
                c.alignment = Alignment(vertical='top', wrap_text=True)

    # Ringkasan
    sr = r0 + len(rows) + 1
    ws.merge_cells(f'A{sr}:{last_col}{sr}')
    ws[f'A{sr}'] = (f"RINGKASAN: Total {len(rows)} pengajuan | Menunggu {t['pending']} | "
                    f"Terverifikasi {t['verified']} | Ditolak {t['rejected']} | "
                    f"Total Qty {t['qty']}")
    ws[f'A{sr}'].font = bold_font

    # Blok TTD dua pihak
    sr2 = sr + 3
    left_c, right_c = 'C', 'G'
    ws[f'{left_c}{sr2}'] = 'Dibuat oleh,'
    ws[f'{right_c}{sr2}'] = 'Mengetahui,'
    for col in (left_c, right_c):
        ws[f'{col}{sr2}'].font = normal_font
    sr3 = sr2 + 5
    ws[f'{left_c}{sr3}'] = fin_name.upper()
    ws[f'{right_c}{sr3}'] = head_name.upper()
    for col in (left_c, right_c):
        ws[f'{col}{sr3}'].font = bold_font
        ws[f'{col}{sr3}'].border = Border(bottom=Side(style='thin'))
    sr4 = sr3 + 1
    ws[f'{left_c}{sr4}'] = 'Finance'
    ws[f'{right_c}{sr4}'] = 'Kepala Cabang'
    for col in (left_c, right_c):
        ws[f'{col}{sr4}'].font = Font(name='Arial', italic=True, size=8, color='64748B')

    # Print setup: landscape, fit to width
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ============================================================
# PDF (fpdf2) — landscape A4, kop identitas perusahaan
# ============================================================
class WaterReportPDF:
    """Wrapper tipis di atas BPFBasePDF (reuse kop/footer/font)."""

    def __init__(self):
        from modules import pdf_generator as _pg
        from modules.pdf_generator import BPFBasePDF
        self._pg = _pg
        self.pdf = BPFBasePDF(orientation='L', unit='mm', format='A4')
        self.pdf.set_auto_page_break(auto=True, margin=18)
        self.pdf.set_margins(10, 7, 10)

    @property
    def INK(self): return self._pg.INK

    @property
    def GRAY_LABEL(self): return self._pg.GRAY_LABEL

    @property
    def RULE(self): return self._pg.RULE

    # kolom: No | No Dokumen | Tanggal | OB | Rincian | Qty | Status | Remark | Catatan
    WIDTHS = [10, 34, 20, 34, 62, 14, 28, 52, 52]
    HEADERS = ['No', 'No. Dokumen', 'Tanggal', 'OB', 'Rincian Item', 'Qty',
               'Status', 'Remark Verifikasi', 'Catatan']

    def generate(self, rows, meta):
        p = self.pdf
        head_name, fin_name = _signatures(meta)
        t = _totals(rows)
        p.add_page()
        self._title(meta)
        self._table(rows)
        self._summary(t)
        self._signatures(head_name, fin_name)
        out = p.output()  # fpdf2 ≥2.7: bytearray
        return bytes(out) if isinstance(out, (bytearray, memoryview)) else str(out).encode('latin-1', errors='replace')

    def _title(self, meta):
        p = self.pdf
        ident = meta.get('company') or {}
        p.set_font(p._font(), 'B', 13)
        p.set_text_color(*self.INK)
        p.cell(0, 7, 'LAPORAN REKAP PENGAJUAN AIR MINUM', align='C',
               new_x='LMARGIN', new_y='NEXT')
        p.set_font(p._font(), '', 8.5)
        p.set_text_color(*self.GRAY_LABEL)
        sub = (f"{ident.get('company_name', 'PT BESTPROFIT FUTURES')} — "
               f"{ident.get('company_subtitle', 'Cabang')}"
               f"  •  Periode {_fmt_date(meta.get('from'))} s/d {_fmt_date(meta.get('to'))}")
        p.cell(0, 5, p.clean_text(sub), align='C', new_x='LMARGIN', new_y='NEXT')
        ftxt = meta.get('filters_text') or 'Semua status'
        p.set_font(p._font(), 'I', 7.5)
        p.cell(0, 4.5, p.clean_text(f'Filter: {ftxt}'), align='C',
               new_x='LMARGIN', new_y='NEXT')
        p.ln(3)
        p.set_text_color(*self.INK)

    def _row_h(self, texts):
        p = self.pdf
        # tinggi baris = max baris wrap (kolom Rincian/Remark/Catatan paling luas)
        lines = 1
        for w, txt in zip(self.WIDTHS, texts):
            if w in (self.WIDTHS[4], self.WIDTHS[7], self.WIDTHS[8]):
                p.set_font(p._font(), '', 7)
                lines = max(lines, max(1, len(p.multi_cell(w - 2, 3, p.clean_text(str(txt or '-')), dry_run=True, output="LINES"))))
        return lines * 3 + 2.5

    def _table(self, rows):
        p = self.pdf
        p.set_draw_color(*self.RULE)
        p.set_fill_color(29, 78, 216)
        p.set_text_color(255, 255, 255)
        p.set_font(p._font(), 'B', 7.5)
        p.set_line_width(0.2)
        for w, h in zip(self.WIDTHS, self.HEADERS):
            p.cell(w, 7, h, border=1, align='C', fill=True)
        p.ln()
        p.set_text_color(*self.INK)
        fill = False
        for i, r in enumerate(rows, 1):
            vals = [str(i), r.get('display_id', '-'), _fmt_date(r.get('purchase_date')),
                    r.get('ob_name', '-'), _items_text(r),
                    str(sum(int(it.get('quantity') or 0) for it in (r.get('items') or []))),
                    STATUS_LABEL.get(r.get('status'), r.get('status', '-')),
                    r.get('remark') or '-', r.get('note') or '-']
            h = self._row_h(vals)
            if p.get_y() + h > p.h - 22:
                self._page_break_continue(rows, i)
                fill = False
            x0 = p.l_margin
            p.set_xy(x0, p.get_y())
            p.set_font(p._font(), '', 7)
            if fill:
                p.set_fill_color(241, 245, 249)
            aligns = ['C', 'L', 'C', 'L', 'L', 'C', 'C', 'L', 'L']
            wrap_idx = {4, 7, 8}
            x = x0
            for ci, (w, v) in enumerate(zip(self.WIDTHS, vals)):
                p.set_xy(x, p.get_y())
                if ci in wrap_idx:
                    p.multi_cell(w, 3, p.clean_text(str(v)), border=1 if not fill else 1,
                                 align=aligns[ci], fill=fill)
                    x += w
                    p.set_xy(x, p.get_y())
                else:
                    p.cell(w, h, p.clean_text(str(v)), border=1, align=aligns[ci], fill=fill)
                    x += w
            p.ln(h)
            fill = not fill

    def _page_break_continue(self, rows, next_i):
        p = self.pdf
        p.add_page()
        self._table_header_again()

    def _table_header_again(self):
        p = self.pdf
        p.set_draw_color(*self.RULE)
        p.set_fill_color(29, 78, 216)
        p.set_text_color(255, 255, 255)
        p.set_font(p._font(), 'B', 7.5)
        for w, h in zip(self.WIDTHS, self.HEADERS):
            p.cell(w, 7, h, border=1, align='C', fill=True)
        p.ln()
        p.set_text_color(*self.INK)

    def _summary(self, t):
        p = self.pdf
        p.ln(2)
        p.set_font(p._font(), 'B', 8)
        p.cell(0, 5, p.clean_text(
            f"RINGKASAN: Total {t['pending'] + t['verified'] + t['rejected']} pengajuan  |  "
            f"Menunggu {t['pending']}  |  Terverifikasi {t['verified']}  |  "
            f"Ditolak {t['rejected']}  |  Total Qty {t['qty']}"),
            new_x='LMARGIN', new_y='NEXT')

    def _signatures(self, head_name, fin_name):
        p = self.pdf
        if p.get_y() + 34 > p.h - 16:
            p.add_page()
        p.ln(6)
        p.set_font(p._font(), '', 8)
        p.set_text_color(*self.INK)
        col_w = 80
        gap = (p.w - p.l_margin - p.r_margin - 2 * col_w) / 2
        y0 = p.get_y()
        # kiri: Finance (Dibuat oleh)
        p.set_xy(p.l_margin, y0)
        p.cell(col_w, 5, 'Dibuat oleh,', align='C')
        # kanan: Kepala Cabang (Mengetahui)
        p.set_xy(p.l_margin + col_w + gap, y0)
        p.cell(col_w, 5, 'Mengetahui,', align='C')
        p.ln(18)
        p.set_draw_color(*self.GRAY_LABEL)
        p.set_line_width(0.3)
        y_line = p.get_y()
        p.line(p.l_margin + 8, y_line, p.l_margin + col_w - 8, y_line)
        p.line(p.l_margin + col_w + gap + 8, y_line, p.l_margin + col_w + gap + col_w - 8, y_line)
        p.ln(2)
        p.set_font(p._font(), 'B', 9)
        p.set_xy(p.l_margin, p.get_y())
        p.cell(col_w, 5, p.clean_text(fin_name).upper(), align='C')
        p.set_xy(p.l_margin + col_w + gap, p.get_y())
        p.cell(col_w, 5, p.clean_text(head_name).upper(), align='C')
        p.ln(5)
        p.set_font(p._font(), 'I', 7)
        p.set_text_color(*self.GRAY_LABEL)
        p.set_xy(p.l_margin, p.get_y())
        p.cell(col_w, 4, 'Finance', align='C')
        p.set_xy(p.l_margin + col_w + gap, p.get_y())
        p.cell(col_w, 4, 'Kepala Cabang', align='C')
