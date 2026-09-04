"""API Routes - Transactions, Cross-Check, Analytics, Stats"""
from flask import request, jsonify, session
from modules.config import get_db_connection
from modules.helpers import log_activity_async, safe_float, role_required
from modules.engine import generate_human_insight
from datetime import datetime, timedelta

def register_transaction_api(app):

    @app.route('/api/transactions/archive')
    @role_required(['ga', 'finance', 'admin'])
    def api_archive_transactions():
        try:
            page = request.args.get('page', 1, type=int)
            limit = request.args.get('limit', 50, type=int)
            search = request.args.get('search', '').strip()
            start_date = request.args.get('start_date', '')
            end_date = request.args.get('end_date', '')
            bbm_type = request.args.get('bbm_type', '')

            conn = get_db_connection()
            if not conn: return jsonify({'error': 'DB error'}), 500
            cursor = conn.cursor(dictionary=True)

            where = ["status = 'archived'"]
            params = []
            if not start_date:
                where.append("created_at >= DATE_SUB(CURDATE(), INTERVAL 7 DAY)")
            if start_date:
                where.append("DATE(created_at) >= %s"); params.append(start_date)
            if end_date:
                where.append("DATE(created_at) <= %s"); params.append(end_date)
            if search:
                where.append("(nopol LIKE %s OR driver_name LIKE %s)")
                params.extend([f"%{search}%", f"%{search}%"])
            if bbm_type:
                where.append("bbm_type = %s"); params.append(bbm_type)

            wc = " AND ".join(where)
            cursor.execute(f"SELECT COUNT(*) as total FROM transactions WHERE {wc}", params)
            total = cursor.fetchone()['total']
            cursor.execute(f"SELECT COALESCE(SUM(nominal),0) as total_nominal FROM transactions WHERE {wc}", params)
            s = cursor.fetchone()
            cursor.execute(f"SELECT * FROM transactions WHERE {wc} ORDER BY created_at DESC LIMIT %s OFFSET %s", params + [limit, (page-1)*limit])
            data = cursor.fetchall()

            for row in data:
                for k in ['nominal','liter','price_per_liter','odo_km','km_per_liter']:
                    if row.get(k) is not None: row[k] = float(row[k])

            cursor.close(); conn.close()
            return jsonify({
                'data': data, 'total': total, 'page': page, 'limit': limit,
                'total_pages': max(1, (total + limit - 1) // limit),
                'summary': {'total_nominal': float(s['total_nominal'])}
            })
        except Exception as e:
            print(f'[ERROR] api_archive_transactions: {e}')
            return jsonify({'error': 'Terjadi kesalahan server'}), 500

    @app.route('/api/stats')
    @role_required(['ga', 'finance', 'admin'])
    def api_stats():
        try:
            conn = get_db_connection()
            if not conn: return jsonify({'error': 'DB error'}), 500
            cursor = conn.cursor(dictionary=True)
            cursor.execute("SELECT COUNT(*) as c FROM transactions"); total = cursor.fetchone()['c']
            cursor.execute("SELECT COUNT(*) as c FROM transactions WHERE status IN ('pending','modified')"); pending = cursor.fetchone()['c']
            cursor.execute("SELECT COUNT(*) as c FROM transactions WHERE status='verified_ga'"); vga = cursor.fetchone()['c']
            cursor.execute("SELECT COUNT(*) as c FROM transactions WHERE status='os_finance'"); osf = cursor.fetchone()['c']
            cursor.execute("SELECT COUNT(*) as c FROM transactions WHERE status='archived'"); arc = cursor.fetchone()['c']
            # Ringkasan hari ini (context bar dashboard)
            cursor.execute("SELECT COUNT(*) as c, COALESCE(SUM(nominal),0) as n FROM transactions WHERE DATE(created_at) = CURDATE()")
            today = cursor.fetchone()
            cursor.execute("SELECT COUNT(*) as c FROM transactions WHERE DATE(created_at) = CURDATE() AND status IN ('pending','modified')"); tp = cursor.fetchone()['c']
            cursor.execute("SELECT COUNT(*) as c FROM transactions WHERE DATE(created_at) = CURDATE() AND status='verified_ga'"); tg = cursor.fetchone()['c']
            cursor.execute("SELECT COUNT(*) as c FROM transactions WHERE DATE(created_at) = CURDATE() AND status='os_finance'"); tos = cursor.fetchone()['c']
            cursor.execute("SELECT COUNT(*) as c FROM transactions WHERE DATE(created_at) = CURDATE() AND status='archived'"); ta = cursor.fetchone()['c']
            cursor.close(); conn.close()
            return jsonify({
                'total': total, 'pending': pending, 'verified_ga': vga, 'os_finance': osf, 'archived': arc,
                'today_tx': today['c'], 'today_nominal': float(today['n'] or 0),
                'today_pending': tp, 'today_verified_ga': tg, 'today_os_finance': tos, 'today_archived': ta,
            })
        except Exception as e:
            print(f'[ERROR] api_stats: {e}')
            return jsonify({'error': 'Terjadi kesalahan server'}), 500

    @app.route('/api/audit-logs')
    @role_required(['ga', 'finance', 'admin'])
    def api_audit_logs():
        """Audit log. Bila ?branch=<code> diberikan (admin), lihat log dari DB cabang itu."""
        try:
            branch = request.args.get('branch', '').strip().upper()
            conn = get_db_connection(branch_code=branch or None)
            if not conn:
                return jsonify({'error': 'DB error'}), 500
            cursor = conn.cursor(dictionary=True)
            cursor.execute("SELECT id, transaction_id, action, user_type, user_name, created_at, ip_address, branch_code FROM activity_logs ORDER BY created_at DESC LIMIT 500")
            logs = cursor.fetchall(); cursor.close(); conn.close()
            return jsonify(logs)
        except Exception as e:
            print(f'[ERROR] api_audit_logs: {e}')
            return jsonify({'error': 'Terjadi kesalahan server'}), 500

    @app.route('/api/cross-check/<int:tx_id>')
    @role_required(['ga', 'finance', 'admin'])
    def api_cross_check(tx_id):
        log_activity_async(tx_id, 'cross_check_view', 'ga', 'GA Officer', new_data={'action': 'view_cross_check'}, ip=request.remote_addr, ua=request.headers.get('User-Agent'))
        try:
            conn = get_db_connection()
            if not conn: return jsonify({'error': 'DB error'}), 500
            cursor = conn.cursor(dictionary=True)
            cursor.execute("SELECT * FROM transactions WHERE id = %s", (tx_id,))
            tx = cursor.fetchone()
            if not tx: cursor.close(); conn.close(); return jsonify({'error': 'Transaction not found'}), 404

            cursor.execute("SELECT odo_km, km_per_liter, created_at FROM transactions WHERE nopol = %s AND status = 'archived' AND id < %s ORDER BY id DESC LIMIT 1", (tx['nopol'], tx_id))
            prev = cursor.fetchone()
            cursor.execute("SELECT ROUND(AVG(NULLIF(km_per_liter,0)), 2) as avg_kml, COUNT(*) as tx_count FROM transactions WHERE nopol = %s AND status = 'archived' AND km_per_liter > 0 AND created_at >= DATE_SUB(NOW(), INTERVAL 3 MONTH)", (tx['nopol'],))
            avg_data = cursor.fetchone()
            cursor.execute("SELECT COUNT(*) as total_tx, COALESCE(SUM(nominal),0) as total_nominal FROM transactions WHERE driver_name = %s AND status = 'archived' AND MONTH(created_at) = MONTH(CURDATE())", (tx['driver_name'],))
            monthly = cursor.fetchone()
            cursor.execute("SELECT config_value FROM system_config WHERE config_key = 'monthly_budget'")
            budget_row = cursor.fetchone(); budget = float(budget_row['config_value']) if budget_row else 3000000

            avg_kml = float(avg_data['avg_kml']) if avg_data and avg_data['avg_kml'] else 10.0
            health_score = min(100, max(0, int((avg_kml / 14) * 100)))
            flags = []
            odo_diff = float(tx['odo_km']) - float(prev['odo_km']) if prev else 0.0
            if odo_diff < 0: flags.append({'level': 'danger', 'msg': 'ODO MUNDUR!'})
            elif odo_diff == 0: flags.append({'level': 'warning', 'msg': 'ODO tidak berubah'})
            elif odo_diff < 30: flags.append({'level': 'warning', 'msg': f'Jarak tempuh hanya {odo_diff} km'})

            current_kml = float(tx['km_per_liter']) if tx['km_per_liter'] and float(tx['km_per_liter']) > 0 else None
            if current_kml and avg_data and avg_data['avg_kml']:
                if current_kml < float(avg_data['avg_kml']) * 0.7: flags.append({'level': 'warning', 'msg': f'KM/L ({current_kml}) di bawah rata-rata ({avg_data["avg_kml"]})'})
                elif current_kml > float(avg_data['avg_kml']) * 1.5: flags.append({'level': 'warning', 'msg': f'KM/L ({current_kml}) di atas rata-rata ({avg_data["avg_kml"]})'})

            budget_usage = (float(monthly['total_nominal']) / budget * 100) if budget > 0 else 0
            if budget_usage > 80: flags.append({'level': 'warning', 'msg': f'Budget {budget_usage:.0f}%'})
            elif budget_usage > 100: flags.append({'level': 'danger', 'msg': f'Budget HABIS! ({budget_usage:.0f}%)'})

            has_danger = any(f['level'] == 'danger' for f in flags)
            overall = 'danger' if has_danger else ('warning' if any(f['level']=='warning' for f in flags) else 'success')
            cursor.close(); conn.close()

            return jsonify({
                'current': {'id': tx['id'], 'nopol': tx['nopol'], 'driver_name': tx['driver_name'], 'odo_km': float(tx['odo_km']), 'nominal': float(tx['nominal']), 'liter': float(tx['liter']), 'km_per_liter': float(tx['km_per_liter']) if tx['km_per_liter'] else 0},
                'previous_odo': {'odo_km': prev['odo_km'], 'km_per_liter': prev['km_per_liter'], 'date': str(prev['created_at'])} if prev else None,
                'odo_diff': odo_diff, 'avg_3months': {'avg_kml': avg_data['avg_kml'], 'tx_count': avg_data['tx_count']} if avg_data else None,
                'monthly': {'total_tx': monthly['total_tx'], 'total_nominal': float(monthly['total_nominal'])},
                'budget': budget, 'budget_usage_percent': round(budget_usage, 1), 'health_score': health_score,
                'flags': flags, 'overall': overall,
                'recommendation': 'AMAN' if overall=='success' else ('PERLU PERHATIAN' if overall=='warning' else 'INVESTIGASI!')
            })
        except Exception as e:
            print(f'[ERROR] api_cross_check: {e}')
            return jsonify({'error': 'Terjadi kesalahan server'}), 500

    @app.route('/api/transaction-flags')
    @role_required(['ga', 'finance', 'admin'])
    def api_transaction_flags():
        try:
            conn = get_db_connection()
            if not conn: return jsonify({'error': 'DB error'}), 500
            cursor = conn.cursor(dictionary=True)
            cursor.execute("SELECT id, nopol, driver_name, odo_km FROM transactions WHERE status IN ('pending','modified') ORDER BY created_at ASC")
            txs = cursor.fetchall()
            result = {}
            for tx in txs:
                flags = []
                cursor.execute("SELECT odo_km FROM transactions WHERE nopol = %s AND status = 'archived' AND id < %s ORDER BY id DESC LIMIT 1", (tx['nopol'], tx['id']))
                prev = cursor.fetchone()
                if prev:
                    odo_diff = float(tx['odo_km']) - float(prev['odo_km'])
                    if odo_diff < 0: flags.append({'level': 'danger', 'msg': 'ODO mundur'})
                    elif odo_diff == 0: flags.append({'level': 'warning', 'msg': 'ODO tidak berubah'})
                    elif odo_diff < 30: flags.append({'level': 'warning', 'msg': f'Top-up ({odo_diff} km)'})
                result[str(tx['id'])] = {'flags': flags, 'overall': 'danger' if any(f['level']=='danger' for f in flags) else ('warning' if any(f['level']=='warning' for f in flags) else 'success')}
            cursor.close(); conn.close()
            return jsonify(result)
        except Exception as e:
            print(f'[ERROR] api_transaction_flags: {e}')
            return jsonify({'error': 'Terjadi kesalahan server'}), 500

    @app.route('/api/vehicle-health')
    @role_required(['ga', 'finance', 'admin'])
    def api_vehicle_health():
        try:
            conn = get_db_connection()
            if not conn: return jsonify({'error': 'DB error'}), 500
            cursor = conn.cursor(dictionary=True)
            cursor.execute("""SELECT t.nopol, t.vehicle_type, COUNT(t.id) as total_tx, ROUND(AVG(NULLIF(t.km_per_liter,0)), 2) as avg_kml, MAX(t.created_at) as last_activity, SUM(t.nominal) as total_nominal, COALESCE(SUM(t.jumlah_appointment),0) as total_appt, (SELECT va.driver_name FROM vehicle_assignments va WHERE va.nopol = t.nopol AND va.is_current = 1 ORDER BY va.id DESC LIMIT 1) as current_driver FROM transactions t WHERE t.status = 'archived' GROUP BY t.nopol, t.vehicle_type ORDER BY avg_kml DESC""")
            units = cursor.fetchall()
            for unit in units:
                kml = unit['avg_kml'] if unit['avg_kml'] else 10
                days_since = 30
                if unit['last_activity']: days_since = (datetime.now() - unit['last_activity']).days
                health = min(100, int(min(100, (kml/14)*60) + max(0, 20-min(20, days_since)) + min(20, unit['total_tx']*2)))
                unit['health_score'] = health
                unit['status'] = 'good' if health >= 70 else ('warning' if health >= 40 else 'danger')
            cursor.close(); conn.close()
            return jsonify({'units': units, 'total_active_units': len(units), 'avg_fleet_health': round(sum(u['health_score'] for u in units)/len(units)) if units else 0})
        except Exception as e:
            print(f'[ERROR] api_vehicle_health: {e}')
            return jsonify({'error': 'Terjadi kesalahan server'}), 500

    @app.route('/api/get-performance/<plat_nomor>')
    @role_required(['ga', 'finance', 'admin'])
    def get_performance(plat_nomor):
        try:
            conn = get_db_connection(); cursor = conn.cursor(dictionary=True)
            cursor.execute("SELECT km_per_liter, jumlah_appointment FROM transactions WHERE nopol=%s AND status IN ('verified','archived','os_finance','verified_ga','rejected') AND km_per_liter>0 AND created_at>=DATE_SUB(NOW(), INTERVAL 1 MONTH)", (plat_nomor,))
            data = cursor.fetchall(); cursor.close(); conn.close()
            if not data: return jsonify({"nopol": plat_nomor, "status": "BELUM CUKUP DATA", "avg_kml": 0})
            avg_kml = sum([d['km_per_liter'] for d in data if d['km_per_liter']]) / len(data)
            total_apt = sum([d['jumlah_appointment'] for d in data if d['jumlah_appointment']])
            return jsonify({"nopol": plat_nomor, "status": "BAIK" if avg_kml>=10 else "BOROS", "avg_kml": round(avg_kml,2), "total_appointment": total_apt})
        except Exception as e:
            print(f'[ERROR] get_performance: {e}')
            return jsonify({"error": 'Terjadi kesalahan server'}), 500

    @app.route('/api/get-feedback/<nopol>')
    @role_required(['ga', 'finance', 'admin'])
    def get_vehicle_feedback(nopol):
        try:
            conn = get_db_connection(); cursor = conn.cursor(dictionary=True)
            cursor.execute("SELECT km_per_liter, jumlah_appointment FROM transactions WHERE nopol=%s AND status IN ('verified','archived','os_finance','verified_ga','rejected') AND km_per_liter>0 ORDER BY created_at DESC LIMIT 10", (nopol,))
            data = cursor.fetchall(); cursor.close(); conn.close()
            if not data: return jsonify({"status": "info", "msg": f"Belum ada data untuk {nopol}."})
            km_values = [d['km_per_liter'] for d in data if d['km_per_liter'] and d['km_per_liter']>0]
            if not km_values: return jsonify({"status": "info", "msg": f"Data KM/L belum lengkap."})
            avg_kpl = sum(km_values)/len(km_values)
            performa = "SANGAT BAIK" if avg_kpl>=12 else ("BAIK" if avg_kpl>=10 else ("CUKUP" if avg_kpl>=8 else "BOROS"))
            msg = generate_human_insight(performa, avg_kpl, 12.0, 0, len(km_values))
            return jsonify({"status": "success", "avg_km_per_liter": round(avg_kpl,2), "performa": performa, "msg": msg})
        except Exception as e:
            print(f'[ERROR] get_vehicle_feedback: {e}')
            return jsonify({"status": "error", "msg": 'Terjadi kesalahan server'}), 500

    @app.route('/api/analytics/data')
    @role_required(['ga', 'finance', 'admin'])
    def api_analytics_data():
        try:
            start_date = request.args.get('start_date', '')
            end_date = request.args.get('end_date', '')
            driver = request.args.get('driver', '').strip()
            nopol = request.args.get('nopol', '').strip()
            tx_type = request.args.get('type', '').strip()

            # Default: 1 bulan terakhir
            if not start_date:
                start_date = (datetime.now() - timedelta(days=30)).strftime('%Y-%m-%d')
            if not end_date:
                end_date = datetime.now().strftime('%Y-%m-%d')

            # Build WHERE clause
            where = ["status='archived'", "DATE(created_at) >= %s", "DATE(created_at) <= %s"]
            params = [start_date, end_date]
            if driver:
                where.append("driver_name = %s")
                params.append(driver.upper())
            if nopol:
                where.append("nopol = %s")
                params.append(nopol.upper())
            if tx_type:
                where.append("transaction_type = %s")
                params.append(tx_type)
            wc = " AND ".join(where)

            conn = get_db_connection()
            if not conn: return jsonify({'error': 'DB error'}), 500
            cursor = conn.cursor(dictionary=True)
            cursor.execute(f"SELECT COALESCE(SUM(nominal),0) as total_month, COUNT(*) as total_tx FROM transactions WHERE {wc}", params)
            fin = cursor.fetchone()
            cursor.execute(f"SELECT DATE_FORMAT(created_at, '%Y-%m') as month, SUM(nominal) as total FROM transactions WHERE {wc} GROUP BY month ORDER BY month", params)
            monthly = cursor.fetchall()
            cursor.execute(f"SELECT driver_name, nopol, SUM(nominal) as total, COUNT(*) as tx_count FROM transactions WHERE {wc} GROUP BY driver_name, nopol ORDER BY total DESC LIMIT 5", params)
            top_drivers = cursor.fetchall()
            cursor.execute("SELECT COUNT(*) as c FROM drivers WHERE is_active=1"); total_drivers = cursor.fetchone()['c']
            cursor.execute(f"SELECT COALESCE(SUM(jumlah_appointment),0) as c FROM transactions WHERE {wc}", params); total_appt = int(cursor.fetchone()['c'])
            cursor.execute(f"SELECT driver_name, COUNT(*) as c FROM transactions WHERE {wc} GROUP BY driver_name ORDER BY c DESC LIMIT 1", params)
            top_driver_row = cursor.fetchone()
            cursor.execute(f"SELECT COUNT(*) as c FROM transactions WHERE {wc}", params); total_claims = cursor.fetchone()['c']
            cursor.execute("SELECT nopol, ROUND(AVG(NULLIF(km_per_liter,0)),2) as avg_kml FROM transactions WHERE status='archived' AND km_per_liter > 0 GROUP BY nopol ORDER BY avg_kml DESC")
            eff = cursor.fetchall()
            cash_total = 0; cash_amount = 0
            if not tx_type or tx_type == 'CASH_LPJ':
                cash_where = "status='COMPLETED'"
                cash_params = []
                if start_date:
                    cash_where += " AND DATE(created_at) >= %s"; cash_params.append(start_date)
                if end_date:
                    cash_where += " AND DATE(created_at) <= %s"; cash_params.append(end_date)
                cursor.execute(f"SELECT COUNT(*) as c FROM fuel_cash_requests WHERE {cash_where}", cash_params)
                cash_total = cursor.fetchone()['c']
                cursor.execute(f"SELECT COALESCE(SUM(total_amount),0) as amount FROM fuel_cash_requests WHERE {cash_where}", cash_params)
                cash_amount = cursor.fetchone()['amount']
            cursor.close(); conn.close()
            return jsonify({
                'finance': {
                    'total_month': float(fin['total_month']), 'total_tx': fin['total_tx'], 'avg_per_day': round(float(fin['total_month'])/30), 'avg_per_tx': round(float(fin['total_month'])/fin['total_tx']) if fin['total_tx'] > 0 else 0,
                    'monthly_labels': [m['month'] for m in monthly], 'monthly_amounts': [float(m['total']) for m in monthly],
                    'top_drivers': [{'driver_name': t['driver_name'], 'nopol': t['nopol'], 'total': float(t['total']), 'tx_count': t['tx_count']} for t in top_drivers]
                },
                'ga': {
                    'total_drivers': total_drivers, 'total_claims': total_claims,
                    'total_appt': total_appt, 'top_driver': top_driver_row['driver_name'] if top_driver_row else '-',
                    'freq_labels': [t['driver_name'] for t in top_drivers],
                    'freq_values': [t['tx_count'] for t in top_drivers],
                    'appt_vs_claim': []
                },
                'cash': {
                    'total': cash_total, 'amount': float(cash_amount),
                    'daily_labels': [], 'daily_values': []
                },
                'fleet': {
                    'best_vehicle': eff[0]['nopol'] if eff else '-',
                    'worst_vehicle': eff[-1]['nopol'] if eff else '-',
                    'avg_kml': round(sum([float(e['avg_kml']) for e in eff])/len(eff),1) if eff else 0,
                    'anomaly_count': 0,
                    'eff_labels': [e['nopol'] for e in eff],
                    'eff_values': [float(e['avg_kml']) for e in eff],
                    'service_alerts': []
                }
            })
        except Exception as e:
            print(f'[ERROR] api_analytics_data: {e}')
            return jsonify({'error': 'Terjadi kesalahan server'}), 500

    @app.route('/api/finance-review/<int:tx_id>')
    @role_required(['finance', 'admin'])
    def api_finance_review(tx_id):
        try:
            conn = get_db_connection()
            if not conn: return jsonify({'error': 'DB error'}), 500
            cursor = conn.cursor(dictionary=True)
            cursor.execute("SELECT * FROM transactions WHERE id = %s", (tx_id,))
            tx = cursor.fetchone()
            if not tx: cursor.close(); conn.close(); return jsonify({'error': 'Not found'}), 404

            # Previous ODO
            cursor.execute("SELECT odo_km, created_at FROM transactions WHERE nopol = %s AND status = 'archived' AND id < %s ORDER BY id DESC LIMIT 1", (tx['nopol'], tx_id))
            prev = cursor.fetchone()

            # Monthly stats
            cursor.execute("SELECT COUNT(*) as total_tx, COALESCE(SUM(nominal),0) as total_nominal FROM transactions WHERE driver_name = %s AND MONTH(created_at) = MONTH(CURDATE())", (tx['driver_name'],))
            monthly = cursor.fetchone()

            # Budget
            cursor.execute("SELECT config_value FROM system_config WHERE config_key = 'monthly_budget'")
            budget_row = cursor.fetchone()
            budget = float(budget_row['config_value']) if budget_row else 3000000

            # Photos
            photos = []
            for field, label in [('foto_odo_sebelum','ODO'),('foto_nota_odo_sesudah','Nota+ODO'),('foto_struk','Struk'),('foto_struk_dispenser','Dispenser')]:
                if tx.get(field): photos.append({'label': label, 'url': '/uploads/' + tx[field]})

            cursor.close(); conn.close()
            return jsonify({
                'transaction': {
                    'id': tx['id'], 'nopol': tx['nopol'], 'driver_name': tx['driver_name'],
                    'vehicle_type': tx.get('vehicle_type', 'AVANZA'), 'bbm_type': tx.get('bbm_type', 'PERTALITE'),
                    'nominal': float(tx['nominal']), 'liter': float(tx['liter']),
                    'price_per_liter': float(tx.get('price_per_liter', 0)),
                    'odo_km': float(tx['odo_km']), 'km_per_liter': float(tx.get('km_per_liter', 0) or 0),
                    'jumlah_appointment': tx.get('jumlah_appointment', 0),
                    'spbu_type': tx.get('spbu_type', 'rekanan'), 'gps_address': tx.get('gps_address', ''),
                    'created_at': str(tx['created_at']), 'status': tx['status']
                },
                'previous_odo': {'odo_km': prev['odo_km'], 'date': str(prev['created_at'])} if prev else None,
                'monthly': {'total_tx': monthly['total_tx'], 'total_nominal': float(monthly['total_nominal'])},
                'budget': budget,
                'photos': photos
            })
        except Exception as e:
            print(f'[ERROR] api_finance_review: {e}')
            return jsonify({'error': 'Terjadi kesalahan server'}), 500

    @app.route('/api/finance-remark', methods=['POST'])
    @role_required(['finance', 'admin'])
    def api_finance_remark():
        try:
            data = request.get_json()
            tx_id = data.get('tx_id'); remark = data.get('remark', '').strip()
            username = session.get('full_name', 'Finance Officer')
            if not tx_id or not remark: return jsonify({'status': 'error', 'msg': 'ID transaksi dan remark wajib'}), 400
            conn = get_db_connection(); cursor = conn.cursor()
            cursor.execute("UPDATE transactions SET transaction_notes = CONCAT(COALESCE(transaction_notes,''), '\n[', NOW(), '] ', %s, ': ', %s) WHERE id = %s", (username, remark, tx_id))
            conn.commit(); cursor.close(); conn.close()
            log_activity_async(tx_id, 'finance_remark', 'finance', username, new_data={'remark': remark}, ip=request.remote_addr)
            return jsonify({'status': 'success', 'msg': 'Remark berhasil disimpan'})
        except Exception as e:
            print(f'[ERROR] api_finance_remark: {e}')
            return jsonify({'status': 'error', 'msg': 'Terjadi kesalahan server'}), 500

    @app.route('/api/trip-detail/<int:trip_id>')
    @role_required(['ga', 'finance', 'admin'])
    def api_trip_detail(trip_id):
        try:
            conn = get_db_connection()
            if not conn: return jsonify({'error': 'DB error'}), 500
            cursor = conn.cursor(dictionary=True)
            cursor.execute("SELECT * FROM trip_masters WHERE id = %s", (trip_id,))
            master = cursor.fetchone()
            if not master: cursor.close(); conn.close(); return jsonify({'error': 'Trip not found'}), 404
            cursor.execute("""SELECT td.*, a.display_id AS appointment_display, a.nasabah_name AS appointment_nasabah
                             FROM trip_details td
                             LEFT JOIN appointments a ON td.appointment_id = a.id
                             WHERE td.trip_master_id = %s ORDER BY td.no_urut""", (trip_id,))
            details = cursor.fetchall(); cursor.close(); conn.close()
            if master.get('jam_keberangkatan'): master['jam_keberangkatan'] = str(master['jam_keberangkatan'])
            if master.get('jam_tiba'): master['jam_tiba'] = str(master['jam_tiba'])
            for d in details:
                if d.get('pukul_berangkat'): d['pukul_berangkat'] = str(d['pukul_berangkat'])
                if d.get('pukul_tujuan'): d['pukul_tujuan'] = str(d['pukul_tujuan'])
            return jsonify({'master': master, 'details': details})
        except Exception as e:
            print(f'[ERROR] api_trip_detail: {e}')
            return jsonify({'error': 'Terjadi kesalahan server'}), 500

    @app.route('/api/queue/export-excel')
    @role_required(['ga', 'finance', 'admin'])
    def api_queue_export_excel():
        """Export antrean transaksi (tab aktif) ke file Excel (v2.21)."""
        try:
            from io import BytesIO
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment

            tab = request.args.get('tab', 'ga')
            status_map = {
                'ga': 'pending',
                'finance': 'verified_ga',
                'confirm': 'os_finance',
                'archive': 'archived',
            }
            status = status_map.get(tab, 'pending')

            conn = get_db_connection()
            if not conn: return jsonify({'error': 'DB error'}), 500
            cursor = conn.cursor(dictionary=True)
            cursor.execute(
                "SELECT id, display_id, created_at, nopol, driver_name, bbm_type, "
                "nominal, liter, odo_km, km_per_liter, status FROM transactions "
                "WHERE status=%s ORDER BY created_at DESC LIMIT 500", (status,))
            rows = cursor.fetchall()
            cursor.close(); conn.close()

            wb = Workbook()
            ws = wb.active
            ws.title = 'Antrean ' + tab.upper()
            headers = ['ID', 'Display', 'Tanggal', 'Nopol', 'Driver', 'BBM',
                       'Nominal', 'Liter', 'ODO', 'KM/L', 'Status']
            ws.append(headers)
            hdr_fill = PatternFill('solid', fgColor='1F2937')
            for c in ws[1]:
                c.font = Font(bold=True, color='FFFFFF')
                c.fill = hdr_fill
                c.alignment = Alignment(horizontal='center')
            for r in rows:
                ws.append([
                    r.get('id'), r.get('display_id'),
                    (r.get('created_at') or '').strftime('%d-%m-%Y %H:%M') if hasattr(r.get('created_at'), 'strftime') else r.get('created_at'),
                    r.get('nopol'), r.get('driver_name'), r.get('bbm_type'),
                    float(r.get('nominal') or 0), float(r.get('liter') or 0),
                    int(r.get('odo_km') or 0), float(r.get('km_per_liter') or 0),
                    r.get('status'),
                ])
            for col in ws.columns:
                letter = col[0].column_letter
                ws.column_dimensions[letter].width = max(12, max((len(str(c.value or '')) for c in col), default=8) + 2)
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=7, max_col=9):
                for c in row:
                    c.number_format = '#,##0'

            buf = BytesIO()
            wb.save(buf)
            buf.seek(0)
            from flask import make_response
            resp = make_response(buf.getvalue())
            resp.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            resp.headers['Content-Disposition'] = f'attachment; filename=antrean_{tab}_{datetime.now().strftime("%Y%m%d_%H%M")}.xlsx'
            return resp
        except Exception as e:
            print(f'[ERROR] api_queue_export_excel: {e}')
            return jsonify({'error': 'Terjadi kesalahan server'}), 500
